from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"D:\云栖码梦\云栖码梦_已有功能概况.docx"
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_run(run, size=11.5, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.75)
sec.bottom_margin = Inches(.7)
sec.left_margin = Inches(.85)
sec.right_margin = Inches(.85)
for style_name, size in [("Normal", 11.5), ("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)]:
    s = doc.styles[style_name]
    s.font.name = "Microsoft YaHei"
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    if style_name != "Normal":
        s.font.bold = True

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("云栖码梦")
set_run(r, 20, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(16)
r = p.add_run("已有功能概况")
set_run(r, 14, True, BLUE)

p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.25
p.paragraph_format.space_after = Pt(12)
r = p.add_run("项目定位：")
set_run(r, 11.5, True, RED)
r = p.add_run("面向老年人的 AI 健康管理与数字生活辅助应用，采用老人端和家属端双角色设计，当前已完成核心移动端页面、后端接口和 Android APK。")
set_run(r)

data = [
    ("一 用户与账号功能", [
        "用户注册与登录，支持老人端和家属端角色。",
        "昵称、头像等个人资料管理。",
        "密码显示与隐藏、密码强度提示和基础输入校验。",
        "JWT 登录认证和本地登录状态保存。"]),
    ("二 老人端功能", [
        "今日用药提醒：查看当天用药安排、服药时间和剂量。",
        "用药录入：支持手动录入药品，也支持拍照录入辅助识别。",
        "服药打卡：记录待服、已服和漏服状态。",
        "AI 问事：咨询手机操作、便民办事、生活常识和健康科普问题。",
        "语音交互：支持语音输入和文字播报。",
        "家人互动：查看已绑定家属并进入对话。",
        "个人中心：管理头像、昵称、字体大小、通知设置和亲属绑定。"]),
    ("三 家属端功能", [
        "健康总览：查看本人或已绑定长辈的健康信息。",
        "长辈选择：支持在多个已绑定长辈之间切换。",
        "用药查看：查看长辈用药计划和用药记录。",
        "体征查看：查看血压、血糖、心率、体重等健康数据和趋势。",
        "报告查看：查看长辈上传的体检报告及相关解读。",
        "代录功能：在授权范围内协助录入用药和健康数据。",
        "家人互动：与已绑定长辈进行信息交流。"]),
    ("四 家属绑定功能", [
        "老人端和家属端均支持生成邀请码。",
        "家属输入邀请码后建立绑定关系。",
        "绑定关系支持列表查看和状态识别。",
        "未绑定时可从家人互动页面直接进入个人中心的亲属绑定管理。",
        "后端接口会校验角色、绑定状态和目标用户，限制越权访问。"]),
    ("五 AI 与识别功能", [
        "基于大语言模型的数字生活和健康科普问答。",
        "药品图片 OCR 识别和药品信息整理。",
        "体检报告图片文字识别和内容解释。",
        "症状自查和基础风险分级提示。",
        "知识库检索，覆盖手机教程、便民服务、生活常识和健康科普。",
        "健康场景中展示免责声明，不将 AI 输出作为诊断或处方。"]),
    ("六 健康数据功能", [
        "血压、血糖、心率和体重等体征记录。",
        "体征历史列表和趋势图展示。",
        "体征数据导出。",
        "问诊录音上传、医院和医生信息记录及回放。",
        "天气信息和健康提醒。"]),
    ("七 当前可验证成果", [
        "已生成 Android Debug APK，可安装到真实手机测试。",
        "已完成 FastAPI 后端服务，可通过 8000 端口提供接口。",
        "已完成手机与电脑后端联调流程。",
        "已验证注册登录、角色进入、家属绑定、用药管理和 AI 问答等核心流程。",
        "已形成启动脚本、测试说明和接口文档，支持后续部署和复用。"]),
]

for title, items in data:
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(9)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(title)
    set_run(r, 15, True)
    for item in items:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.first_line_indent = Inches(-.18)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run("• ")
        set_run(r, 11.5, True, RED)
        r = p.add_run(item)
        set_run(r)

h = doc.add_paragraph(style="Heading 1")
h.paragraph_format.space_before = Pt(12)
r = h.add_run("功能状态说明")
set_run(r, 15, True)
p = doc.add_paragraph()
p.paragraph_format.line_spacing = 1.25
p.paragraph_format.space_after = Pt(0)
r = p.add_run("当前版本已具备")
set_run(r, 11.5, True, RED)
r = p.add_run("可运行、可演示和可验证的核心功能。部分高级能力，例如稳定的公网部署、正式消息推送、复杂医疗场景专业审核和大规模并发，还需要在后续版本中继续完善。")
set_run(r)

doc.save(OUT)
print(OUT)
