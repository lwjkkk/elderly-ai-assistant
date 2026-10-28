
import os as _archive_os
if _archive_os.environ.get('RUN_ARCHIVED_ONEOFF') != '1':
    raise SystemExit('Historical one-off script. Review paths before explicitly enabling RUN_ARCHIVED_ONEOFF=1')
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"D:\云栖码梦\云栖码梦_开源赋能AI应用创新_十分钟答辩稿.docx"
RED = RGBColor(0xC0, 0x00, 0x00)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def add_para(doc, text="", style=None, space_after=6, first_indent=True):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(21)
    p.add_run(text)
    return p

def add_rich(doc, parts, style=None, space_after=6, first_indent=True):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(21)
    for text, red in parts:
        r = p.add_run(text)
        if red:
            r.bold = True
            r.font.color.rgb = RED
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11.5)
for name, size in [("Title", 20), ("Heading 1", 15), ("Heading 2", 12.5)]:
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
p.add_run("云栖码梦").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("开源赋能的 AI 应用创新赛道 十分钟答辩稿")
r.bold = True
r.font.size = Pt(13)
r.font.color.rgb = BLUE

add_rich(doc, [("使用说明：", True), ("红色加粗文字是答辩时需要重点强调的关键词。全文按十分钟演讲节奏编排，可直接朗读。", False)], first_indent=False, space_after=10)

sections = [
("一 开场与项目定位", [
("尊敬的各位评委老师，大家好。我是“云栖码梦——银龄数字生活辅助智能体”项目的答辩人。我们选择的参赛方向是“开源赋能的 AI 应用创新”。", []),
("“云栖码梦”不是一个单纯的聊天机器人，也不是只展示概念的页面原型，而是一个面向老年人真实生活场景的 AI 应用系统。我们的目标是让老年人更容易使用智能设备，更容易理解健康信息，也让家属能够在获得授权后参与老人的健康管理。", [("开源赋能的 AI 应用创新", True), ("老人端、家属端、后端服务和 Android APK", True)]),
("目前项目已经完成注册登录、老人和家属角色、家属绑定、用药记录、健康数据管理和 AI 问答等核心流程，可以进行实际运行和演示。接下来我将从需求场景、产品实现、开源融合、技术架构、实证效果和开源复用六个方面进行介绍。", [("实际运行和演示", True)]),
]),
("二 真实问题与目标用户", [
("我们关注的是老年人在数字生活中的两个具体困难：第一是不会使用智能手机，第二是健康信息管理困难。老人可能不会网上挂号、看不懂药品说明、记不清服药时间，也不知道如何记录和理解血压、血糖等数据。", [("真实问题", True), ("老年人", True)]),
("很多现有应用功能很多，但是面向年轻人设计，存在字体小、按钮小、操作路径长和专业术语多等问题。因此，我们的目标用户包括需要大字体、大按钮和语音辅助的老年用户，以及希望远程了解和协助老人健康管理的家属用户。", [("适老化交互", True), ("老人端和家属端", True)]),
("我们的产品定位是健康科普、信息整理和数字生活辅助工具，不是医疗诊断工具，也不替代医生进行诊疗和处方。", [("健康科普、信息整理和数字生活辅助工具", True), ("不替代医生", True)]),
]),
("三 产品功能与应用场景", [
("第一个核心场景是用药管理。老人可以手动录入药品，也可以通过拍照辅助识别药品信息。系统支持记录药名、剂量、服药频次、开始使用药品日期和具体服药时间点，并生成今日用药安排。老人服药后可以进行打卡，家属绑定后也可以查看老人的用药记录。", [("用药管理", True), ("拍照辅助识别", True), ("今日用药安排", True)]),
("第二个场景是拍照查药和体检报告解读。系统通过 OCR 识别图片中的文字，再利用 AI 对药品或报告内容进行整理，用更容易理解的语言帮助用户阅读。系统不直接输出疾病诊断，而是提供信息解释和风险提示，并提醒用户遵医嘱。", [("OCR", True), ("AI 信息整理", True), ("不直接输出疾病诊断", True)]),
("第三个场景是体征记录。用户可以记录血压、血糖、心率和体重，系统通过列表和趋势图帮助用户观察变化，并支持导出，方便就医时提供参考。", [("体征记录", True), ("趋势图", True)]),
("第四个场景是家属绑定与健康协同。老人生成邀请码，家属输入邀请码建立绑定关系。绑定后，家属可以查看授权范围内的用药、体征和报告信息，也可以协助录入。", [("家属绑定与健康协同", True), ("授权范围", True)]),
("第五个场景是 AI 问事和语音交互。老人可以询问手机操作、医保挂号、防诈骗和生活常识等问题。系统支持文字输入、语音输入和语音播报，把复杂操作转换成一步一步的大白话说明。", [("AI 问事和语音交互", True), ("一步一步的大白话说明", True)]),
]),
("四 开源技术选型与融合方式", [
("本项目使用 Flutter 构建移动端，使用 FastAPI 构建后端服务，使用 SQLAlchemy 和 SQLite 进行数据管理，并结合 OCR、Qwen、LangChain、ChromaDB、APScheduler、fl_chart、speech_to_text 和 flutter_tts 等开源技术。", [("Flutter", True), ("FastAPI", True), ("Qwen", True), ("LangChain", True), ("ChromaDB", True)]),
("我们没有把开源组件停留在简单调用层面，而是围绕老年人的真实需求进行了二次整合。例如，将 OCR 识别结果接入药品和报告解释流程，将知识库检索接入数字生活问答，将模型回答限制在健康科普和信息整理范围内，并将老人端、家属端和绑定权限组织成完整业务流程。", [("二次整合", True), ("完整业务流程", True)]),
("因此，项目的核心不是某一个单独的模型或框架，而是开源技术组合后形成的可运行应用。", [("开源技术组合", True), ("可运行应用", True)]),
]),
("五 技术架构与安全设计", [
("项目采用前后端分离架构。Flutter 前端负责页面展示、角色切换、图片上传、语音交互和数据可视化；FastAPI 后端负责认证、绑定、用药、体征、报告、AI 问答和数据同步。", [("前后端分离架构", True)]),
("数据访问采用角色和绑定关系双重控制。老人端默认访问自己的数据；家属端必须是家属角色，并且与目标老人存在有效绑定关系，才能查看对应数据。不能仅凭一个老人 ID 访问其他用户的信息。", [("角色和绑定关系双重控制", True), ("不能越权访问", True)]),
("在健康安全方面，系统加入免责声明和风险提示。我们明确产品是健康科普和信息整理助手，不替代医生诊断；对于高风险和紧急情况，引导用户联系家属、咨询医生或及时就医。", [("健康科普和信息整理助手", True), ("不替代医生诊断", True), ("风险提示", True)]),
]),
("六 创新点与赛道匹配", [
("本项目的第一项创新，是开源技术面向真实人群进行应用创新。我们不是为了展示模型而做模型，而是把 OCR、大语言模型、知识库和移动端能力落到老年人的用药、健康记录和手机操作场景中。", [("真实人群", True), ("应用创新", True)]),
("第二项创新，是将健康管理和数字生活辅助结合。老人既可以问“这个药怎么理解”，也可以问“手机怎么操作”，并且可以在同一个应用中完成记录、查询和管理。", [("健康管理和数字生活辅助结合", True)]),
("第三项创新，是老人端和家属端形成协同闭环。老人拥有自己的账号和数据，家属需要先绑定，再在授权范围内查看和协助，兼顾了老人的独立使用和家庭照护。", [("协同闭环", True), ("独立使用和家庭照护", True)]),
("第四项创新，是对开源能力进行适老化改造，包括大字号、大按钮、简化页面、语音交互、大白话提示和健康免责声明。", [("适老化改造", True), ("大字号、大按钮、语音交互", True)]),
]),
("七 实证成果与现场演示", [
("目前项目已经完成可以运行的 Android APK、后端服务和真机测试流程。现场演示可以按照注册登录、用药录入、用药打卡、AI 问答、家属绑定、健康总览和体征趋势的顺序进行。", [("Android APK", True), ("真机测试流程", True)]),
("我们重点展示三个可验证结果：第一，核心功能能够实际运行；第二，老人端和家属端之间可以形成数据协同；第三，开源组件确实参与了核心功能，而不是只出现在技术说明中。", [("可验证结果", True), ("核心功能能够实际运行", True), ("数据协同", True)]),
]),
("八 开源成果与复用计划", [
("按照开源创新赛道的要求，我们计划整理并提交项目源代码、开源使用清单、知识库说明、部署文档、使用说明、演示视频和验证材料。开源使用清单会明确记录项目名称、版本号、来源、许可证类型、使用方式、修改内容和开源边界。", [("开源使用清单", True), ("版本号、来源、许可证类型", True)]),
("后续我们会继续完善 README、LICENSE、NOTICE、部署教程和二次开发说明，形成其他团队可以下载、运行和修改的项目。我们希望项目不仅是一个参赛作品，也能成为其他开发者开展老年人数字服务时的参考。", [("可下载、可运行、可二次开发", True), ("复用价值", True)]),
]),
("九 不足与未来规划", [
("目前项目仍处于持续完善阶段。AI 回答质量需要通过更多真实问题进行测试，语音识别还需要提升对方言、噪声和老年人语速的适应能力，系统也需要进一步加强并发能力、数据安全和消息推送。", [("持续完善", True), ("真实问题测试", True)]),
("未来我们计划增加离线缓存、异常数据提醒、更加细致的隐私控制和更完善的开源社区维护机制，让项目从可运行原型逐步发展为可复用、可推广的银龄数字生活辅助工具。", [("可复用、可推广", True), ("开源社区维护机制", True)]),
]),
("十 结束语", [
("最后，我想用一句话总结我们的项目：我们不是把开源技术简单放进一个 App，而是用开源技术解决老年人真实的数字生活问题。", [("不是简单拼接开源技术", True), ("解决真实问题", True)]),
("“云栖码梦”希望让老人少一点面对智能手机时的无助，多一点独立使用数字服务的能力；也希望让家属能够在尊重老人隐私的前提下，提供更加及时的陪伴和帮助。", [("尊重隐私", True), ("及时陪伴和帮助", True)]),
("我的汇报到此结束，谢谢各位评委老师。", []),
]),
]

for title, paras in sections:
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(5)
    h.add_run(title)
    for text, reds in paras:
        reds = [item[0] if isinstance(item, tuple) else item for item in reds]
        parts = []
        pos = 0
        for key in reds:
            idx = text.find(key, pos)
            if idx < 0:
                continue
            if idx > pos:
                parts.append((text[pos:idx], False))
            parts.append((key, True))
            pos = idx + len(key)
        if pos < len(text):
            parts.append((text[pos:], False))
        add_rich(doc, parts, first_indent=True)

doc.add_paragraph()
h = doc.add_paragraph(style="Heading 1")
h.add_run("答辩时必须强调的关键词")
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "关键词"
hdr[1].text = "建议表达"
for c in hdr:
    set_cell_shading(c, "1F4E79")
    for p in c.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(255,255,255)
            r.bold = True
for key, val in [
    ("开源融合", "我们不是简单调用模型，而是将多个开源组件组合成完整业务流程。"),
    ("实际验证", "项目已经完成 Android APK、后端服务和真机测试，可以现场演示核心流程。"),
    ("复用价值", "我们会提供源代码、开源使用清单、许可证说明、部署文档和二次开发说明。"),
    ("安全合规", "产品定位是健康科普和信息整理助手，不替代医生诊断，并通过权限和免责声明控制风险。"),
]:
    cells = table.add_row().cells
    cells[0].text = key
    cells[1].text = val
    for r in cells[0].paragraphs[0].runs:
        r.font.color.rgb = RED
        r.bold = True

doc.add_paragraph()
add_para(doc, "提示：红色加粗内容是本赛道评分重点，答辩时要主动说出，尤其是“开源融合”“实际验证”“复用价值”和“安全合规”。", first_indent=False, space_after=0)
doc.save(OUT)
print(OUT)
