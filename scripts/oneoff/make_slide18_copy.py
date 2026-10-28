
import os as _archive_os
if _archive_os.environ.get('RUN_ARCHIVED_ONEOFF') != '1':
    raise SystemExit('Historical one-off script. Review paths before explicitly enabling RUN_ARCHIVED_ONEOFF=1')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = r"D:\云栖码梦\云栖码梦答辩_第18页_创新点修改版.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

NAVY = RGBColor(31, 68, 119)
TEXT = RGBColor(22, 22, 22)
MUTED = RGBColor(74, 74, 74)
LIGHT = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)

def box(x, y, w, h, fill=None, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill or WHITE
    shape.line.color.rgb = line or fill or WHITE
    shape.line.width = Pt(1.2)
    return shape

def text(x, y, w, h, value, size=18, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, font="Microsoft YaHei", margin=0.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape

# Header
box(0, 0, 13.333, 1.05, NAVY, NAVY)
tabs = ["需求与场景", "方案架构", "AI 与开源融合", "自主实现内容", "功能效果与创新", "开放成果与后续计划"]
widths = [2.2, 2.15, 2.35, 2.2, 2.15, 2.283]
x = 0
for i, (label, w) in enumerate(zip(tabs, widths)):
    active = i == 4
    if active:
        s = box(x, 0, w, 1.05, WHITE, WHITE, True)
        s.adjustments[0] = 0.18
        text(x + .05, .18, w - .1, .62, label, 21, NAVY, True, PP_ALIGN.CENTER)
    else:
        text(x + .04, .25, w - .08, .42, label, 18, WHITE, False, PP_ALIGN.CENTER)
    x += w

text(.62, 1.32, 12.0, .48, "三个核心创新：让开源 AI 真正服务老年人", 25, TEXT, True)

# Left takeaway
box(.62, 2.02, 3.55, 4.7, LIGHT, NAVY, True)
text(.93, 2.35, 2.95, .35, "核心价值", 19, NAVY, True)
text(.93, 2.95, 2.95, 1.42, "从“能回答”\n到“能理解、能操作、能协同”", 26, TEXT, True)
text(.93, 4.72, 2.95, 1.2, "Qwen 本地推理\nRAG 知识库\nPaddleOCR\nFlutter + FastAPI", 17, MUTED)
text(.93, 6.03, 2.95, .4, "技术支撑，而非创新本身", 14, NAVY, True)

# Right innovations
box(4.55, 2.02, 8.15, 4.7, WHITE, NAVY, True)
items = [
    ("① 老人—家属协同闭环",
     "邀请码绑定 + elder_id 数据隔离。老人自主记录，家属授权查看、代录和提醒。",
     "从单人健康记录扩展为家庭协同照护"),
    ("② 开源 AI 适老化改造",
     "Qwen 本地推理 + RAG + OCR + 语音交互，配合大字体、大按钮和步骤化回答。",
     "把开源能力转化为老人能执行的服务"),
    ("③ 健康信息安全闭环",
     "OCR 识别 → AI 整理 → 风险提示 → 家属协同 → 就医引导。",
     "兼顾应用效果与医疗场景安全边界"),
]
y = 2.35
for title, body, highlight in items:
    text(4.9, y, 7.35, .35, title, 20, NAVY, True)
    text(4.9, y + .42, 7.35, .56, body, 16.5, TEXT)
    text(4.9, y + 1.02, 7.35, .3, "亮点：" + highlight, 15.5, RGBColor(170, 0, 0), True)
    y += 1.43

prs.save(OUT)
print(OUT)
